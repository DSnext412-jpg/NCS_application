"""Report exporters: CSV, Excel (openpyxl) and PDF (reportlab).

All three consume the same :class:`app.services.report_service.ReportData`
model so the on-screen tables, CSV, Excel and PDF always show identical
numbers. Files are written into ``Paths().reports_dir``.
"""

from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Callable

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.paths import Paths
from app.services.report_service import ReportData, report_filename

logger = logging.getLogger(__name__)

_LINE = colors.HexColor("#334155")
_MUTED = colors.HexColor("#64748b")
_DARK = colors.HexColor("#0f172a")
_ACCENT = colors.HexColor("#2563eb")


def reports_dir() -> Path:
    path = Paths().reports_dir
    path.mkdir(parents=True, exist_ok=True)
    return path


def default_export_path(report: ReportData, ext: str) -> Path:
    """Deterministic default location for a report export."""
    return reports_dir() / report_filename(report.report_type, report.date_from, report.date_to, ext)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def export_csv(report: ReportData, path: str | Path | None = None) -> Path:
    """Write a UTF-8 CSV; one file per report, tables stacked with blank rows."""
    target = Path(path) if path else default_export_path(report, "csv")
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([report.hotel_name])
        writer.writerow([report.title])
        writer.writerow([f"Period: {report.date_from} to {report.date_to}"])
        writer.writerow([f"Generated: {report.generated_at:%Y-%m-%d %H:%M}"])
        writer.writerow([])
        for label, value in report.summary:
            writer.writerow([label, value])
        writer.writerow([])
        for table in report.tables:
            writer.writerow([table.name])
            writer.writerow(table.headers)
            writer.writerows(table.rows)
            writer.writerow([])
    return target


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def export_excel(report: ReportData, path: str | Path | None = None) -> Path:
    """Write an .xlsx workbook; summary on the first sheet, one sheet per table."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    target = Path(path) if path else default_export_path(report, "xlsx")
    target.parent.mkdir(parents=True, exist_ok=True)

    title_font = Font(bold=True, size=14)
    sub_font = Font(size=10, color="595959")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    center = Alignment(horizontal="left", vertical="center")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append([report.hotel_name])
    sheet["A1"].font = title_font
    sheet.append([report.title])
    sheet["A2"].font = sub_font
    sheet.append([f"Period: {report.date_from} to {report.date_to}"])
    sheet.append([f"Generated: {report.generated_at:%Y-%m-%d %H:%M}"])
    sheet.append([])
    for label, value in report.summary:
        sheet.append([label, value])
        sheet[f"A{sheet.max_row}"].font = Font(bold=True)
    sheet.append([])
    sheet.append(["Tables", "Rows"])
    sheet[f"A{sheet.max_row}"].font = header_font
    sheet[f"B{sheet.max_row}"].font = header_font
    for table in report.tables:
        sheet.append([table.name, len(table.rows)])
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width, 60)

    for table in report.tables:
        ws = workbook.create_sheet(title=table.name[:31] or "Sheet")
        ws.append(table.headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for row in table.rows:
            ws.append(row)
        for column_cells in ws.columns:
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width, 40)

    workbook.save(target)
    return target


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _pdf_styles() -> dict[str, ParagraphStyle]:
    return {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=_DARK),
        "sub": ParagraphStyle("sub", fontName="Helvetica", fontSize=9, leading=13, textColor=_MUTED),
        "section": ParagraphStyle(
            "section", fontName="Helvetica-Bold", fontSize=11, leading=15, textColor=_ACCENT,
            spaceBefore=6, spaceAfter=4,
        ),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=8.5, leading=11, textColor=_DARK),
        "cell_bold": ParagraphStyle("cell_bold", fontName="Helvetica-Bold", fontSize=8.5, leading=11, textColor=_DARK),
    }


def export_pdf(report: ReportData, path: str | Path | None = None) -> Path:
    """Write an A4 PDF with the report title, summary and its tables."""
    target = Path(path) if path else default_export_path(report, "pdf")
    target.parent.mkdir(parents=True, exist_ok=True)

    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(target),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=report.title,
        author=report.hotel_name,
    )
    story = [
        Paragraph(report.hotel_name, styles["title"]),
        Spacer(1, 2),
        Paragraph(report.title, styles["sub"]),
        Paragraph(
            f"Period: {report.date_from} to {report.date_to}  |  "
            f"Generated: {report.generated_at:%Y-%m-%d %H:%M}",
            styles["sub"],
        ),
        Spacer(1, 8),
    ]

    if report.summary:
        summary_rows = [[Paragraph(label, styles["cell_bold"]), Paragraph(value, styles["cell"])] for label, value in report.summary]
        summary_table = Table(summary_rows, colWidths=[62 * mm, 92 * mm])
        summary_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(summary_table)
        story.append(Spacer(1, 10))

    for table in report.tables:
        story.append(Paragraph(table.name, styles["section"]))
        header_row = [Paragraph(header, styles["cell_bold"]) for header in table.headers]
        body = [[Paragraph(cell, styles["cell"]) for cell in row] for row in table.rows]
        data = [header_row] + body
        table_obj = Table(data, hAlign="LEFT", repeatRows=1)
        table_obj.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table_obj)
        story.append(Spacer(1, 8))

    doc.build(story)
    return target


EXPORTERS: dict[str, Callable[[ReportData, str | Path | None], Path]] = {
    "csv": export_csv,
    "xlsx": export_excel,
    "pdf": export_pdf,
}


def export_report(report: ReportData, fmt: str, path: str | Path | None = None) -> Path:
    """Export any report in ``csv`` / ``xlsx`` / ``pdf`` format."""
    exporter = EXPORTERS.get(fmt.lower())
    if exporter is None:
        raise ValueError(f"Unsupported export format: {fmt!r}")
    return exporter(report, path)
