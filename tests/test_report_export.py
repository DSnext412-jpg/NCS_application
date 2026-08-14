"""Tests for the CSV / Excel / PDF report exporters.

Each exporter writes to an explicit temp path (never the production
``reports`` directory) and produces a real, readable file.
"""

from __future__ import annotations

from datetime import date

import pytest

from app.database.database import Database
from app.services import report_export, report_service
from app.services.settings_service import initialize_default_settings
from tests._db_support import copy_db_template


def _make_database(tmp_path):  # noqa: ANN001
    database = Database(copy_db_template(tmp_path / "hotel.db"))
    return database


def _sample_report(database) -> report_service.ReportData:  # noqa: ANN001
    with database.session_scope() as session:
        return report_service.build_report(
            session,
            "financial",
            date.today(),
            date.today(),
        )


def test_export_csv_writes_readable_file(tmp_path) -> None:  # noqa: ANN001
    database = _make_database(tmp_path)
    try:
        with database.session_scope() as session:
            initialize_default_settings(session)
        report = _sample_report(database)
        target = report_export.export_csv(report, tmp_path / "out.csv")
        assert target.exists()
        text = target.read_text(encoding="utf-8-sig")
        assert "Financial Summary" in text
        assert "Collected" in text
    finally:
        database.dispose()


def test_export_excel_writes_workbook(tmp_path) -> None:  # noqa: ANN001
    pytest.importorskip("openpyxl")
    from openpyxl import load_workbook

    database = _make_database(tmp_path)
    try:
        with database.session_scope() as session:
            initialize_default_settings(session)
        report = _sample_report(database)
        target = report_export.export_excel(report, tmp_path / "out.xlsx")
        assert target.exists()
        workbook = load_workbook(target)
        assert "Summary" in workbook.sheetnames
        assert "Payment Methods" in workbook.sheetnames
    finally:
        database.dispose()


def test_export_pdf_writes_pdf(tmp_path) -> None:  # noqa: ANN001
    database = _make_database(tmp_path)
    try:
        with database.session_scope() as session:
            initialize_default_settings(session)
        report = _sample_report(database)
        target = report_export.export_pdf(report, tmp_path / "out.pdf")
        assert target.exists()
        assert target.stat().st_size > 500
        assert target.read_bytes()[:4] == b"%PDF"
    finally:
        database.dispose()


def test_export_report_dispatch(tmp_path) -> None:  # noqa: ANN001
    database = _make_database(tmp_path)
    try:
        with database.session_scope() as session:
            initialize_default_settings(session)
        report = _sample_report(database)
        csv_path = report_export.export_report(report, "csv", tmp_path / "a.csv")
        pdf_path = report_export.export_report(report, "PDF", tmp_path / "a.pdf")
        assert csv_path.suffix == ".csv"
        assert pdf_path.suffix == ".pdf"
        with pytest.raises(ValueError):
            report_export.export_report(report, "docx", tmp_path / "a.docx")
    finally:
        database.dispose()


def test_export_creates_parent_directories(tmp_path) -> None:  # noqa: ANN001
    database = _make_database(tmp_path)
    try:
        with database.session_scope() as session:
            initialize_default_settings(session)
        report = _sample_report(database)
        deep = tmp_path / "nested" / "deeper" / "out.csv"
        target = report_export.export_csv(report, deep)
        assert target.parent.exists()
    finally:
        database.dispose()
